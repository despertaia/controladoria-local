from datetime import date
from io import BytesIO

import openpyxl

from nucleo import banco
from nucleo.exportar_carteira import COLUNAS, gerar_xlsx


def _planilha():
    conn = banco.conectar(":memory:")
    conn.execute(
        """INSERT INTO processo (numero, tribunal, instancia, cliente, polo_cliente,
               parte_contraria, orgao_julgador, classe, valor_causa, advogado_atua,
               ultimo_andamento_data, ultimo_andamento_texto, ultima_publicacao_data, fontes)
           VALUES ('00000010220248110041', 'TJMT', '1grau', 'MARIA CLIENTE', 'Polo Passivo',
               'BANCO ALFA S.A.', '3ª VARA CÍVEL', '7', 'R$ 1.500,50', 1, '20260917',
               'Juntada de Certidão', '20260910', '["djen", "sistema:cache"]')""")
    conn.execute(
        """INSERT INTO processo (numero, tribunal, advogado_atua, ultima_publicacao_data, fontes)
           VALUES ('00000030220244013600', 'TRF1', NULL, '20260801', '["djen"]')""")
    ativa = conn.execute("SELECT * FROM processo WHERE advogado_atua = 1").fetchall()
    conferir = conn.execute("SELECT * FROM processo WHERE advogado_atua IS NULL").fetchall()
    return openpyxl.load_workbook(BytesIO(gerar_xlsx(ativa, conferir, hoje=date(2026, 10, 5))))


def test_abas_e_cabecalho():
    wb = _planilha()
    assert wb.sheetnames == ["Carteira", "A conferir"]
    assert [c.value for c in wb["Carteira"][1]] == COLUNAS


def test_linhas_formatadas():
    wb = _planilha()
    assert [c.value for c in wb["Carteira"][2]] == [
        "0000001-02.2024.8.11.0041", "TJMT", "1º grau", "MARIA CLIENTE", "Polo Passivo",
        "BANCO ALFA S.A.", "3ª VARA CÍVEL", "7", "R$ 1.500,50", "17/09/2026",
        "Juntada de Certidão", 18, "10/09/2026", "djen, sistema:cache", "Ativo"]
    conferir = [c.value for c in wb["A conferir"][2]]
    assert conferir[0] == "0000003-02.2024.4.01.3600"
    assert conferir[11] == 65
    assert conferir[14] == "A conferir (fora do TJMT)"


def test_dias_sem_movimento_usa_a_data_mais_recente():
    conn = banco.conectar(":memory:")
    conn.execute(
        """INSERT INTO processo (numero, tribunal, instancia, advogado_atua,
               ultimo_andamento_data, ultima_publicacao_data)
           VALUES ('00000010220248110041', 'TJMT', '1grau', 1, '20260101', '20260930')""")
    ativa = conn.execute("SELECT * FROM processo").fetchall()
    wb = openpyxl.load_workbook(BytesIO(gerar_xlsx(ativa, [], hoje=date(2026, 10, 5))))
    assert wb["Carteira"][2][11].value == 5
