"""
Exporta a carteira de processos para uma planilha .xlsx — uma linha por
processo, com partes, advogados, último andamento e intimação pendente — para
o usuário ver tudo de uma vez sem abrir um por um.

Lê do cache (processos já sincronizados); usa openpyxl (já é dependência).
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from captura.processo_parser import formatar_numero_cnj

_COLUNAS = [
    ("Número", 24), ("Classe", 14), ("Órgão / Vara", 28),
    ("Polo Ativo", 30), ("Advogados (ativo)", 30),
    ("Polo Passivo", 30), ("Advogados (passivo)", 30),
    ("Andamentos", 12), ("Último andamento", 14), ("Resumo do último andamento", 50),
    ("Peças", 8), ("Intimação pendente", 16), ("Atualizado em", 18),
]


def _polo(partes, chave: str) -> tuple[str, str]:
    """Retorna (nomes, advogados) do polo cujo rótulo contém `chave`."""
    nomes: list[str] = []
    advs: list[str] = []
    for polo in partes or []:
        if chave.lower() not in str(polo.get("polo", "")).lower():
            continue
        nomes.extend(polo.get("nomes", []))
        for integ in polo.get("integrantes", []):
            for a in integ.get("advogados", []):
                s = a["nome"] + (f' ({a["oab"]})' if a.get("oab") else "")
                advs.append(s)
    return "; ".join(nomes), "; ".join(dict.fromkeys(advs))  # advs sem duplicar


def _intimacao_pendente(partes) -> str:
    for polo in partes or []:
        for integ in polo.get("integrantes", []):
            if integ.get("intimacao_pendente"):
                return "Sim"
    return "Não"


def gerar_planilha_xlsx(registros: list[tuple[dict, dict | None]]) -> bytes:
    """registros: lista de (linha_da_planilha, dados_do_cache_ou_None)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Processos"

    cabecalho_fill = PatternFill("solid", fgColor="1E40AF")
    cabecalho_font = Font(color="FFFFFF", bold=True)
    for c, (titulo, largura) in enumerate(_COLUNAS, start=1):
        cel = ws.cell(row=1, column=c, value=titulo)
        cel.fill = cabecalho_fill
        cel.font = cabecalho_font
        ws.column_dimensions[get_column_letter(c)].width = largura

    for p, d in registros:
        digitos = p.get("numero_digitos") or p.get("numero", "")
        if d:
            numero = d.get("numero_formatado") or formatar_numero_cnj(digitos)
            ativo_n, ativo_a = _polo(d.get("partes"), "Ativo")
            passivo_n, passivo_a = _polo(d.get("partes"), "Passivo")
            ands = d.get("andamentos") or []
            ult = ands[0] if ands else {}
            ws.append([
                numero, d.get("classe", ""), d.get("orgao_julgador", ""),
                ativo_n, ativo_a, passivo_n, passivo_a,
                len(ands), ult.get("data", ""), ult.get("texto", ""),
                len(d.get("documentos") or []), _intimacao_pendente(d.get("partes")),
                (d.get("atualizado_em", "") or "").replace("T", " ")[:16],
            ])
        else:
            # processo ainda não sincronizado: só o que veio da planilha
            ws.append([
                formatar_numero_cnj(digitos), p.get("tipo", ""), p.get("vara", ""),
                p.get("partes", ""), "", "", "",
                "", "", "(não sincronizado)", "", "", "",
            ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUNAS))}{ws.max_row}"
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            cel.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
